"""xlsx, lido pelo openpyxl em modo read_only (streaming)."""

import re
from collections.abc import Iterator, Sequence
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from ingestion.converters.base_converter import FileConverter, Source, batches_from_rows
from ingestion.converters.converter_errors import ConversionError
from ingestion.converters.converter_registry import ConverterFactory

BATCH_ROWS = 10_000


@ConverterFactory.register("xlsx", extensions=(".xlsx", ".xlsm"))
class XlsxConverter(FileConverter):
    """Cada aba vira um Parquet; o cabeçalho está em `config.header_row` (1-based).

    Abre com `data_only=True`: célula com fórmula entrega o valor que o Excel
    calculou, não a fórmula. O xlsx já guarda tipo, então o valor vira texto pela
    regra fixa de `_text` (int → "7", float → repr, data → ISO 8601, bool →
    true/false, vazio → nulo). Duas passadas por aba: a primeira acha a última
    coluna com algum valor, para que colunas vazias de formatação à direita não
    virem schema; coluna sem nome com dado fica, como `column_<n>`.

    Sem `config.sheet`, lê todas as abas (filtradas por `config.include`), com o
    nome da aba como sufixo quando há mais de uma; aba vazia é pulada.
    """

    def _read(self, path: Path) -> Iterator[Source]:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            names = self._sheet_names(path, workbook.sheetnames)
            for name in names:
                sheet = workbook[name]
                sheet.reset_dimensions()  # não confiar na dimensão gravada no arquivo
                width = self._width(sheet)
                if width == 0:
                    continue
                header = next(self._rows(sheet, width), [])
                yield Source(
                    suffix=name if len(names) > 1 else None,
                    header=header,
                    batches=batches_from_rows(
                        self._rows(sheet, width, skip_header=True),
                        width=width,
                        batch_rows=BATCH_ROWS,
                    ),
                )
        finally:
            workbook.close()

    def _sheet_names(self, path: Path, available: Sequence[str]) -> list[str]:
        if self.config.sheet is not None:
            if self.config.sheet not in available:
                raise ConversionError(
                    f"{path.name}: aba {self.config.sheet!r} não existe; "
                    f"abas: {', '.join(repr(a) for a in available)}"
                )
            return [self.config.sheet]
        if self.config.include:
            pattern = re.compile(self.config.include)
            return [name for name in available if pattern.search(name)]
        return list(available)

    def _width(self, sheet: Any) -> int:  # ReadOnlyWorksheet (stubs só cobrem Worksheet)
        width = 0
        for row in sheet.iter_rows(min_row=self.config.header_row, values_only=True):
            for position in range(len(row), width, -1):
                if row[position - 1] is not None:
                    width = position
                    break
        return width

    def _rows(
        self, sheet: Any, width: int, skip_header: bool = False
    ) -> Iterator[list[str | None]]:
        first = self.config.header_row + (1 if skip_header else 0)
        for row in sheet.iter_rows(min_row=first, max_col=width, values_only=True):
            yield [_text(value) for value in row]


def _text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        return repr(value)
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    return str(value)
